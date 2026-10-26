"""
controllers/excel_utils.py
------------------------------
Shared utilities to generate downloadable Excel (.xlsx) files:
export of all the data of a table and empty templates with the
headers the app recognizes on import.

Used by the catalog (books) and by members, so both modules export
with exactly the same style and so that the downloaded file can be
imported again without changes.
"""

import re
import datetime
import unicodedata
from io import BytesIO

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_BLUE = "1D4ED8"


def normalize_header(text):
    """'Título ' -> 'titulo', 'Editorial (con errores)' -> 'editorialconerrores'.

    Removes accents, spaces and symbols so Excel headers written in any
    way can be compared against the names the app recognizes.
    """
    text = str(text or "").strip().lower()
    text = "".join(
        c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-z0-9]", "", text)


def cell_text(value):
    """Converts a cell value into clean text ('' if empty).
    Prevents whole numbers from coming out as '123.0'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def build_excel(sheet_name, headers, rows, widths=None, note=None):
    """Builds an .xlsx file in memory and returns its BytesIO.

    - `headers`: list of column titles (row 1, formatted).
    - `rows`: list of lists/tuples with the values of each row.
    - `widths`: optional list with the width of each column.
    - `note`: optional text added in a second sheet
      ("Instrucciones") to guide whoever fills in the template.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31] or "Datos"

    sheet.append(list(headers))

    fill_style = PatternFill("solid", fgColor=HEADER_BLUE)
    header_font = Font(bold=True, color="FFFFFF")
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.fill = fill_style
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(list(row))

    for position, header in enumerate(headers, start=1):
        if widths and position <= len(widths):
            width = widths[position - 1]
        else:
            width = max(14, min(40, len(str(header)) + 6))
        sheet.column_dimensions[get_column_letter(position)].width = width

    # Freeze the header row and enable the autofilter.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    if note:
        note_sheet = workbook.create_sheet("Instrucciones")
        note_sheet.column_dimensions["A"].width = 110
        for number, line in enumerate(note.strip().splitlines(), start=1):
            cell = note_sheet.cell(row=number, column=1, value=line.strip())
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if number == 1:
                cell.font = Font(bold=True, size=12)

    memory_buffer = BytesIO()
    workbook.save(memory_buffer)
    memory_buffer.seek(0)
    return memory_buffer


def respond_excel(memory_buffer, file_name):
    """Returns the Excel file as a browser download."""
    return send_file(
        memory_buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=file_name,
    )


def name_with_date(base):
    """'miembros' -> 'miembros_2026-09-15.xlsx'."""
    return f"{base}_{datetime.date.today().isoformat()}.xlsx"
